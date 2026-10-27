"""PowerPoint export — pitch-ready valuation deck.

Builds a 4-slide deck from the same data the Excel export uses:
  1. Title
  2. Football Field (native editable bar chart) + valuation range
  3. Trading Comps — peer median multiples
  4. Methodology summary (Comps / DCF / Precedent / WACC)

Design language mirrors the web app: institutional, one accent (deep green),
hairline rules instead of heavy colour bands, square corners, one language.

Native pptx charts/tables → fully editable in PowerPoint, copy-paste into a
master pitch deck.

Custom master / branding
------------------------
Drop a file at ``assets/template.pptx`` and its slide master (background, logo,
theme colours, fonts) is used as the base for every export — the content below is
drawn on top. That is the single place to control the layout that is always used:
edit the master once in PowerPoint and every future deck inherits it. Without a
template the built-in clean styling below applies.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

# ---- Palette (mirrors assets/style.css: ink / graphite / rule / accent green) ----
INK = RGBColor(0x14, 0x18, 0x1D)
GRAPHITE = RGBColor(0x46, 0x4C, 0x55)
MUTED = RGBColor(0x76, 0x7C, 0x86)
RULE = RGBColor(0xC7, 0xCB, 0xD1)
SURFACE = RGBColor(0xF6, 0xF7, 0xF8)
ACCENT = RGBColor(0x0F, 0x5C, 0x4B)        # deep green
ACCENT_SOFT = RGBColor(0xE7, 0xEF, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# One place to control the deck font. IBM Plex Sans matches the web app; if it is
# not installed PowerPoint substitutes a default sans. Change here to restyle all text.
FONT = "IBM Plex Sans"
FONT_MONO = "IBM Plex Mono"

SCALE = 1_000_000  # display in millions

# Consistent geometry
MARGIN = 0.7          # inches, left/right content margin
TITLE_TOP = 0.5
RULE_Y = 1.15         # y of the hairline under a slide title


def _set_run(run, text, size, *, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _textbox(slide, left, top, width, height, text, size=18, *,
             bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size, bold=bold, color=color, font=font)
    return box


def _hline(slide, prs, y, *, x0=MARGIN, x1=None, color=RULE, weight=0.75):
    """Thin horizontal rule — the structural device, in place of colour bands."""
    if x1 is None:
        x1 = prs.slide_width.inches - MARGIN
    ln = slide.shapes.add_connector(2, Inches(x0), Inches(y), Inches(x1), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def _accent_tick(slide, y, *, x=MARGIN, w=0.42, weight=2.4):
    """Short green accent segment above a title — the single accent, used sparingly."""
    ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = ACCENT
    ln.line.width = Pt(weight)
    return ln


def _slide_header(slide, prs, title, eyebrow=None):
    """Consistent content-slide header: accent tick, title, hairline beneath."""
    _accent_tick(slide, TITLE_TOP - 0.06)
    if eyebrow:
        _textbox(slide, MARGIN, TITLE_TOP - 0.02, 10, 0.3, eyebrow.upper(), 10.5,
                 color=MUTED, font=FONT_MONO)
        _textbox(slide, MARGIN, TITLE_TOP + 0.24, 11.9, 0.7, title, 24, bold=True, color=INK)
    else:
        _textbox(slide, MARGIN, TITLE_TOP + 0.02, 11.9, 0.7, title, 24, bold=True, color=INK)
    _hline(slide, prs, RULE_Y, color=RULE)


def _base_presentation() -> Presentation:
    """Load assets/template.pptx if present (uses its master/theme), else a blank deck."""
    tmpl = Path(__file__).resolve().parent / "assets" / "template.pptx"
    prs = Presentation(str(tmpl)) if tmpl.exists() else Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    return prs


def build_pptx(out_path: Path, *, target_name: str, sector: str, ccy: str,
               summary_df, implied: list[dict] | None,
               wacc_result: dict | None, dcf_result: dict | None,
               precedent: dict | None, n_peers_included: int) -> Path:
    prs = _base_presentation()
    blank = prs.slide_layouts[6]
    unit = f"m {ccy}"
    W = prs.slide_width.inches
    # The universe string may carry a leading auto-screen marker ("🔍 ") that is fine
    # in the app UI but reads as clutter on a title slide.
    sector = (sector or "").lstrip("🔍").strip(" ·")

    # ---------------------------------------------------------------- Slide 1: Title
    s = prs.slides.add_slide(blank)
    _accent_tick(s, 2.55, x=MARGIN, w=0.6, weight=3)
    _textbox(s, MARGIN, 2.7, W - 2 * MARGIN, 1.1, "Valuation Summary", 42, bold=True, color=INK)
    _textbox(s, MARGIN, 3.75, W - 2 * MARGIN, 0.6, f"{target_name}   ·   {sector}", 18,
             color=GRAPHITE)
    _hline(s, prs, 4.5, x0=MARGIN, x1=W - MARGIN)
    _textbox(s, MARGIN, 4.62, W - 2 * MARGIN, 0.5,
             "Trading Comps   ·   DCF   ·   Precedent Transactions", 12,
             color=MUTED, font=FONT_MONO)
    _textbox(s, MARGIN, 6.7, W - 2 * MARGIN, 0.4,
             "For discussion purposes only. Not investment advice.", 10, color=MUTED)

    # ---------------------------------------------------------------- Slide 2: Football Field
    s = prs.slides.add_slide(blank)
    _slide_header(s, prs, "Football Field", eyebrow="Implied enterprise value")

    methods = []
    for item in (implied or []):
        if item.get("base_ev"):
            methods.append((item["label"], (item.get("low_ev") or item["base_ev"]) / SCALE,
                            item["base_ev"] / SCALE, (item.get("high_ev") or item["base_ev"]) / SCALE))
    if dcf_result and dcf_result.get("enterprise_value"):
        ev = dcf_result["enterprise_value"] / SCALE
        methods.append(("DCF", ev * 0.9, ev, ev * 1.1))
    if precedent and precedent.get("implied"):
        pick = next((t for t in precedent["implied"] if t["multiple"] == "ev_ebitda"),
                    precedent["implied"][0] if precedent["implied"] else None)
        if pick and pick.get("base_ev"):
            methods.append(("Precedent Tx", (pick.get("low_ev") or pick["base_ev"]) / SCALE,
                            pick["base_ev"] / SCALE, (pick.get("high_ev") or pick["base_ev"]) / SCALE))

    if methods:
        cd = CategoryChartData()
        cd.categories = [m[0] for m in methods]
        cd.add_series("base", [round(m[1], 0) for m in methods])
        cd.add_series("range", [round(m[3] - m[1], 0) for m in methods])
        gframe = s.shapes.add_chart(
            XL_CHART_TYPE.BAR_STACKED, Inches(MARGIN), Inches(1.5),
            Inches(8.1), Inches(5.2), cd)
        chart = gframe.chart
        chart.has_legend = False
        chart.has_title = False
        plot = chart.plots[0]
        plot.gap_width = 70
        ped = chart.series[0]
        ped.format.fill.background()
        ped.format.line.fill.background()
        rng = chart.series[1]
        rng.format.fill.solid()
        rng.format.fill.fore_color.rgb = ACCENT
        rng.format.line.color.rgb = ACCENT
        try:
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.minimum_scale = 0
            for ax in (chart.value_axis, chart.category_axis):
                ax.tick_labels.font.size = Pt(11)
                ax.tick_labels.font.name = FONT
                ax.tick_labels.font.color.rgb = GRAPHITE
        except Exception:
            pass

        # Value callout + per-method table on the right
        rng_lows = [m[1] for m in methods]
        rng_highs = [m[3] for m in methods]
        rx = 9.2
        _textbox(s, rx, 1.5, 3.5, 0.35, "RANGE ACROSS METHODS", 10.5, color=MUTED, font=FONT_MONO)
        _textbox(s, rx, 1.8, 3.6, 0.6,
                 f"{min(rng_lows):,.0f} – {max(rng_highs):,.0f}", 22, bold=True, color=ACCENT)
        _textbox(s, rx, 2.42, 3.6, 0.3, unit, 11, color=MUTED, font=FONT_MONO)
        _hline(s, prs, 2.85, x0=rx, x1=W - MARGIN)

        tbox = s.shapes.add_textbox(Inches(rx), Inches(2.95), Inches(3.7), Inches(3.6))
        tf = tbox.text_frame
        tf.word_wrap = True
        first = True
        for label, low, med, high in methods:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            _set_run(p.add_run(), label, 12.5, bold=True, color=INK)
            p2 = tf.add_paragraph()
            _set_run(p2.add_run(), f"{low:,.0f}  /  {med:,.0f}  /  {high:,.0f}", 11.5,
                     color=GRAPHITE, font=FONT_MONO)
        _textbox(s, rx, 6.55, 3.7, 0.35, f"Low / Median / High  ({unit})", 10, color=MUTED)

    # ---------------------------------------------------------------- Slide 3: Trading Comps
    s = prs.slides.add_slide(blank)
    _slide_header(s, prs, "Trading Comps", eyebrow=f"Peer median  ·  n = {n_peers_included}")

    metric_labels = {"ev_sales": "EV / Sales", "ev_ebitda": "EV / EBITDA",
                     "ev_ebit": "EV / EBIT", "pe": "P / E",
                     "ebitda_margin": "EBITDA margin", "ebit_margin": "EBIT margin"}
    comps_rows = []
    for _, row in summary_df.iterrows():
        m = row["metric"]
        comps_rows.append((metric_labels.get(m, m), row.get("median"), row.get("q1"),
                           row.get("q3"), m.endswith("_margin")))

    n = len(comps_rows) + 1
    tbl_shape = s.shapes.add_table(n, 4, Inches(MARGIN), Inches(1.5),
                                   Inches(7.4), Inches(0.42 * n)).table
    tbl_shape.first_row = False
    tbl_shape.horz_banding = False
    for j, h in enumerate(["Metric", "Median", "Q1", "Q3"]):
        c = tbl_shape.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = ACCENT
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = c.text_frame.paragraphs[0]
        r.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        _set_run(r.add_run(), h, 12.5, bold=True, color=WHITE)
    for i, (label, med, q1, q3, is_margin) in enumerate(comps_rows, start=1):
        fm = (lambda v: f"{v:.1%}" if v is not None else "—") if is_margin \
            else (lambda v: f"{v:.1f}x" if v is not None else "—")
        for j, val in enumerate([label, fm(med), fm(q1), fm(q3)]):
            c = tbl_shape.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else SURFACE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            r = c.text_frame.paragraphs[0]
            r.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            _set_run(r.add_run(), str(val), 11.5, bold=(j == 1),
                     color=INK if j <= 1 else GRAPHITE,
                     font=FONT if j == 0 else FONT_MONO)
    _textbox(s, MARGIN, 1.5 + 0.42 * n + 0.15, 8, 0.35,
             "Outliers excluded from the statistics (IQR method).", 10.5, color=MUTED)

    # ---------------------------------------------------------------- Slide 4: Methodology
    s = prs.slides.add_slide(blank)
    _slide_header(s, prs, "Methodology & Assumptions")

    lines = []
    if implied:
        ev_eb = next((i for i in implied if i["multiple"] == "ev_ebitda"), None)
        if ev_eb:
            lines.append(("Trading Comps",
                          f"Peer median EV/EBITDA {ev_eb['median']:.1f}x  →  implied EV "
                          f"{ev_eb['base_ev']/SCALE:,.0f} {unit}"))
    if wacc_result:
        wl = (f"{wacc_result['wacc']:.2%}  (relevered beta {wacc_result['relevered_beta']:.2f}, "
              f"Ke {wacc_result['cost_of_equity']:.2%})")
        if wacc_result.get("n_de_outliers"):
            wl += f"  ·  {wacc_result['n_de_outliers']} peer(s) with D/E >200% checked against a cleaned beta"
        lines.append(("WACC", wl))
    if dcf_result and dcf_result.get("enterprise_value"):
        lines.append(("DCF",
                      f"Enterprise value {dcf_result['enterprise_value']/SCALE:,.0f} {unit}  "
                      f"(terminal value {dcf_result.get('tv_pct_of_ev', 0):.0%} of EV)"))
    if precedent and precedent.get("stats", {}).get("ev_ebitda"):
        pe = precedent["stats"]["ev_ebitda"]
        lines.append(("Precedent Transactions",
                      f"Deal median EV/EBITDA {pe['median']:.1f}x  ({pe['n']} deals, control premium embedded)"))

    y = 1.55
    for head, body in lines:
        _accent_tick(s, y + 0.07, x=MARGIN, w=0.28, weight=2)
        _textbox(s, MARGIN + 0.42, y - 0.05, 3.3, 0.4, head, 13.5, bold=True, color=INK)
        _textbox(s, MARGIN + 3.9, y - 0.05, W - MARGIN - 4.6, 0.7, body, 12.5, color=GRAPHITE)
        y += 0.95

    _hline(s, prs, 6.5, x0=MARGIN, x1=W - MARGIN)
    _textbox(s, MARGIN, 6.6, W - 2 * MARGIN, 0.5,
             "Data: Yahoo Finance (trading) · CSV/adapter (transactions). "
             "Prototype. For discussion purposes only.", 10, color=MUTED)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
